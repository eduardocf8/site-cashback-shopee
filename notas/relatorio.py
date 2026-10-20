"""Lê o "Relatório Mensal" da Central do Afiliado da Shopee (CSV ou XLSX).

Cada linha é uma loja que pagou comissão extra no mês e vira uma nota. Colunas do
arquivo padrão (outubro/2026):
Mês de conclusão, Nome da loja, ID da Loja, Comissão Total do Vendedor, CNPJ do Vendedor,
CPF do Vendedor, Identificação Fiscal Estrangeira, Razão social do vendedor, Endereço do
Vendedor, País do Vendedor, Inscrição Estadual, E-mail

Também aceita as variações conhecidas: a de agência/MCN ("Taxa Total da MCN" no lugar da
comissão, "E-mail de contato") e a antiga, sem as colunas de estrangeiro."""

import csv
import io
import zipfile
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal, InvalidOperation

from .receita import cnpj_valido

TAMANHO_MAXIMO = 25 * 1024 * 1024
# Muito acima de qualquer relatório real (um afiliado grande tem ~1.000 lojas por mês).
MAXIMO_LINHAS = 20_000
MAXIMO_COLUNAS = 80
MAXIMO_DESCOMPACTADO = 150 * 1024 * 1024  # planilha .xlsx é um zip
MESES = {
    "jan": 1, "feb": 2, "fev": 2, "mar": 3, "apr": 4, "abr": 4, "may": 5, "mai": 5, "jun": 6, "jul": 7,
    "aug": 8, "ago": 8, "sep": 9, "set": 9, "oct": 10, "out": 10, "nov": 11, "dec": 12, "dez": 12,
}

# nome da coluna (sem acento, minúsculo) -> campo
COLUNAS = {
    "mes de conclusao": "mes",
    "nome da loja": "loja_nome",
    "id da loja": "loja_id",
    "comissao total do vendedor": "valor",
    "taxa total da mcn": "valor",
    "cnpj do vendedor": "cnpj",
    "cpf do vendedor": "cpf",
    "identificacao fiscal estrangeira": "nif",
    "razao social do vendedor": "nome",
    "endereco do vendedor": "endereco",
    "pais do vendedor": "pais",
    "e-mail": "email",
    "e-mail de contato": "email",
    "email": "email",
}
OBRIGATORIAS = {"mes", "loja_id", "valor", "nome"}


class ErroRelatorio(ValueError):
    pass


@dataclass
class Linha:
    numero: int  # linha do arquivo (o cabeçalho é a 1)
    mes: date | None
    loja_id: str
    loja_nome: str
    valor: Decimal | None
    tipo_documento: str  # cnpj, cpf, nif ou ""
    documento: str
    nome: str
    email: str
    pais: str
    endereco: str = ""
    problemas: list = field(default_factory=list)

    @property
    def estrangeira(self) -> bool:
        return self.tipo_documento == "nif" or (self.pais and self.pais != "BR")


def _chave(texto: str) -> str:
    sem_acento = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return " ".join(sem_acento.lower().replace("﻿", "").split())


def ler_mes(texto: str) -> date | None:
    t = _chave(texto)
    achado = re.fullmatch(r"([a-z]{3})[a-z]*\.?\s*(?:de\s*)?/?\s*(\d{4})", t)
    if achado and achado.group(1) in MESES:
        return date(int(achado.group(2)), MESES[achado.group(1)], 1)
    achado = re.fullmatch(r"(\d{1,2})[/-](\d{4})", t) or re.fullmatch(r"(\d{4})[/-](\d{1,2})", t)
    if achado:
        a, b = achado.groups()
        ano, mes = (int(b), int(a)) if len(b) == 4 else (int(a), int(b))
        if 1 <= mes <= 12:
            return date(ano, mes, 1)
    return None


def ler_valor(texto) -> Decimal | None:
    if isinstance(texto, (int, float, Decimal)):
        return Decimal(str(texto)).quantize(Decimal("0.01"))
    t = re.sub(r"[^\d,.\-]", "", str(texto or ""))
    if not t:
        return None
    if "," in t:  # formato brasileiro: 1.234,56
        t = t.replace(".", "").replace(",", ".")
    try:
        return Decimal(t).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def _so_digitos(texto) -> str:
    return re.sub(r"\D", "", str(texto or ""))


def _linhas_brutas(dados: bytes, nome_arquivo: str) -> list[list]:
    if nome_arquivo.lower().endswith((".xlsx", ".xlsm")) or dados[:2] == b"PK":
        from openpyxl import load_workbook

        try:
            with zipfile.ZipFile(io.BytesIO(dados)) as zip_:
                descompactado = sum(info.file_size for info in zip_.infolist())
        except zipfile.BadZipFile:
            raise ErroRelatorio("Não consegui abrir essa planilha. Envie o arquivo do jeito que baixou da Shopee.")
        if descompactado > MAXIMO_DESCOMPACTADO:
            raise ErroRelatorio("Essa planilha é grande demais. Envie o relatório de um mês só, do jeito que baixou da Shopee.")
        try:
            planilha = load_workbook(io.BytesIO(dados), read_only=True, data_only=True).worksheets[0]
            # O tamanho que a planilha "diz" ter não vale: um arquivo de 2 KB pode declarar
            # 1 milhão de linhas x 16 mil colunas e esgotar a memória do servidor.
            planilha.reset_dimensions()
            linhas = []
            for linha in planilha.iter_rows(max_col=MAXIMO_COLUNAS, values_only=True):
                if len(linhas) >= MAXIMO_LINHAS:
                    raise ErroRelatorio(f"O relatório tem mais de {MAXIMO_LINHAS:,} linhas.".replace(",", "."))
                linhas.append(["" if c is None else c for c in linha])
        except ErroRelatorio:
            raise
        except Exception:  # noqa: BLE001 - arquivo corrompido ou não é planilha
            raise ErroRelatorio("Não consegui abrir essa planilha. Envie o arquivo do jeito que baixou da Shopee.")
        return linhas
    for codificacao in ("utf-8-sig", "cp1252"):
        try:
            texto = dados.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ErroRelatorio("Não consegui ler esse arquivo. Envie o CSV do jeito que baixou da Shopee.")
    try:
        dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t")
    except csv.Error:
        dialeto = csv.excel
    linhas = []
    for linha in csv.reader(io.StringIO(texto), dialeto):
        if len(linhas) >= MAXIMO_LINHAS:
            raise ErroRelatorio(f"O relatório tem mais de {MAXIMO_LINHAS:,} linhas.".replace(",", "."))
        linhas.append(linha[:MAXIMO_COLUNAS])
    return linhas


def ler(dados: bytes, nome_arquivo: str = "") -> list[Linha]:
    if not dados:
        raise ErroRelatorio("O arquivo está vazio.")
    if len(dados) > TAMANHO_MAXIMO:
        raise ErroRelatorio("Arquivo grande demais (até 25 MB).")
    brutas = [l for l in _linhas_brutas(dados, nome_arquivo) if any(str(c).strip() for c in l)]
    if not brutas:
        raise ErroRelatorio("O arquivo está vazio.")
    indices = {}
    for i, coluna in enumerate(brutas[0]):
        campo = COLUNAS.get(_chave(coluna))
        if campo and campo not in indices:
            indices[campo] = i
    faltando = OBRIGATORIAS - set(indices)
    if faltando:
        raise ErroRelatorio("Esse arquivo não parece o Relatório Mensal da Shopee (Central do Afiliado › "
                            "Relatórios › Relatório mensal › Exportar).")

    def valor(linha, campo):
        i = indices.get(campo)
        return str(linha[i]).strip() if i is not None and i < len(linha) and linha[i] is not None else ""

    saida = []
    for numero, bruta in enumerate(brutas[1:], start=2):
        cnpj, cpf, nif = _so_digitos(valor(bruta, "cnpj")), _so_digitos(valor(bruta, "cpf")), valor(bruta, "nif")
        if cnpj:
            tipo, documento = "cnpj", cnpj.zfill(14)
        elif cpf:
            tipo, documento = "cpf", cpf.zfill(11)
        elif nif:
            tipo, documento = "nif", nif[:40]
        else:
            tipo, documento = "", ""
        linha = Linha(
            numero=numero, mes=ler_mes(valor(bruta, "mes")), loja_id=valor(bruta, "loja_id")[:30],
            loja_nome=valor(bruta, "loja_nome")[:200], valor=ler_valor(valor(bruta, "valor")),
            tipo_documento=tipo, documento=documento, nome=" ".join(valor(bruta, "nome").split())[:300],
            email=valor(bruta, "email").lower()[:254], pais=valor(bruta, "pais").upper()[:2],
            endereco=" ".join(valor(bruta, "endereco").split())[:300],
        )
        if linha.mes is None:
            linha.problemas.append("mês de conclusão ilegível")
        if not linha.loja_id:
            linha.problemas.append("sem ID da loja")
        if linha.valor is None or linha.valor <= 0:
            linha.problemas.append("valor da comissão inválido")
        if not linha.nome:
            linha.problemas.append("sem razão social")
        if not documento:
            linha.problemas.append("sem CNPJ, CPF ou identificação estrangeira")
        elif tipo == "cnpj" and not cnpj_valido(documento):
            linha.problemas.append("CNPJ inválido")
        elif tipo == "cpf" and len(documento) != 11:
            linha.problemas.append("CPF com tamanho errado")
        if linha.email and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", linha.email):
            linha.email = ""  # e-mail quebrado: a nota sai, só não vai por e-mail
        saida.append(linha)
    return saida
