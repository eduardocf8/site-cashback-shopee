"""Planilha do Excel com as vendas do período (aba Resultados > Exportar Excel)."""

import io

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

FORMATO_REAIS = '"R$" #,##0.00'
FORMATO_DATA = "dd/mm/yyyy hh:mm"
CABECALHO = Font(bold=True, color="FFFFFF")
FUNDO_CABECALHO = PatternFill("solid", fgColor="1F3A2E")


def _local(momento):
    return timezone.localtime(momento).replace(tzinfo=None) if momento else None


def _aba(planilha, titulo, colunas, linhas):
    """colunas: [(nome, largura, formato)]. write_only não deixa mudar depois, então o
    cabeçalho e as larguras vão antes das linhas."""
    aba = planilha.create_sheet(titulo)
    for i, (_, largura, _) in enumerate(colunas, start=1):
        aba.column_dimensions[get_column_letter(i)].width = largura
    aba.freeze_panes = "A2"
    cabecalho = []
    for nome, _, _ in colunas:
        celula = _celula(aba, nome)
        celula.font, celula.fill = CABECALHO, FUNDO_CABECALHO
        cabecalho.append(celula)
    aba.append(cabecalho)
    for linha in linhas:
        valores = []
        for valor, (_, _, formato) in zip(linha, colunas):
            celula = _celula(aba, valor)
            if formato and valor is not None:
                celula.number_format = formato
            valores.append(celula)
        aba.append(valores)


def _celula(aba, valor):
    from openpyxl.cell import WriteOnlyCell

    return WriteOnlyCell(aba, value=valor)


def planilha_de_vendas(vendas, produtos, nome_do_grupo) -> bytes:
    planilha = Workbook(write_only=True)
    situacoes = dict(vendas.model.SITUACOES)
    _aba(planilha, "Vendas", [
        ("Data da compra", 17, FORMATO_DATA), ("Pedido", 20, None), ("Produto", 60, None),
        ("ID do produto", 14, None), ("Loja", 28, None), ("Quantidade", 11, "0"),
        ("Preço", 12, FORMATO_REAIS), ("Valor", 12, FORMATO_REAIS), ("Comissão", 12, FORMATO_REAIS),
        ("Status", 12, None), ("Status na Shopee", 16, None), ("Sub ID", 24, None), ("Grupo", 20, None),
        ("Validada em", 17, FORMATO_DATA),
    ], (
        (_local(v.comprado_em), v.pedido_id, v.nome, v.item_id, v.loja, v.quantidade, v.preco, v.valor,
         v.comissao, situacoes.get(v.situacao, v.situacao), v.status_shopee,
         "" if v.sub_id == "sem_subid" else v.sub_id, nome_do_grupo(v.sub_id), _local(v.concluido_em))
        for v in vendas.iterator(chunk_size=2000)
    ))
    _aba(planilha, "Produtos", [
        ("Produto", 60, None), ("ID do produto", 14, None), ("Loja", 28, None), ("Unidades", 10, "0"),
        ("Faturamento", 14, FORMATO_REAIS), ("Comissão", 14, FORMATO_REAIS),
    ], ((p["nome"], p["item_id"], p["loja"], p["unidades"], p["faturamento"], p["comissao"]) for p in produtos))
    saida = io.BytesIO()
    planilha.save(saida)
    return saida.getvalue()
