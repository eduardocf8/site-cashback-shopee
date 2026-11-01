import io
from datetime import date, datetime, timedelta
from decimal import Decimal
from unittest import mock

from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook

from licencas.models import ContaAppfiliado

from . import cripto, resultados
from .models import RegraSubId, VendaItem
from .shopee import ShopeeAPIError, buscar_conversoes
from .tests import SP, BasePainel, conversao, pedido, ts


def item(nome, preco, qtd=1, comissao="0", validado_em=None, item_id="", loja="Loja X", imagem=""):
    return {"itemId": item_id or nome[:10], "itemName": nome, "shopName": loja, "imageUrl": imagem,
            "itemPrice": preco, "qty": qtd, "itemTotalCommission": comissao, "completeTime": validado_em}


class TestGuardarVendas(TestCase):
    AGORA = datetime(2026, 10, 18, 15, 0, tzinfo=SP)

    def setUp(self):
        self.conta = ContaAppfiliado.objects.create(
            email="bia@exemplo.com", shopee_app_id="123", shopee_secret_cifrado=cripto.cifrar("segredo"),
            shopee_validado_em=timezone.now(),
        )

    def sincronizar(self, conversoes, agora=None):
        with mock.patch("painel.resultados.paginas_de_conversoes", return_value=[conversoes]):
            resultados.sincronizar_conta(self.conta, agora=agora or self.AGORA)

    def test_guarda_cada_item_com_status_sub_id_e_loja(self):
        self.sincronizar([
            conversao(ts(2026, 10, 10), [
                pedido("p1", "COMPLETED", [item("Fone TWS", "50,00", 2, "8,00", ts(2026, 10, 15), loja="Loja do Fone"),
                                           item("Capinha", "10", 1, "1")]),
                pedido("p2", "CANCELLED", [item("Mop", "60", 1, "5")]),
            ], utm="grupoA-story---", conversion_id="c1"),
            conversao(ts(2026, 10, 11), [pedido("p3", "PENDING", [item("Garrafa", "40", 1, "4")])], conversion_id="c2"),
        ])
        itens = {v.nome: v for v in VendaItem.objects.filter(conta=self.conta)}
        self.assertEqual(set(itens), {"Fone TWS", "Capinha", "Mop", "Garrafa"})
        fone = itens["Fone TWS"]
        self.assertEqual((fone.pedido_id, fone.loja, fone.quantidade, fone.valor, fone.comissao),
                         ("p1", "Loja do Fone", 2, Decimal("100.00"), Decimal("8.00")))
        self.assertEqual((fone.situacao, fone.sub_id), ("concluido", "grupoA-story"))
        self.assertEqual(timezone.localtime(fone.concluido_em, SP).date().day, 15)
        self.assertEqual(itens["Mop"].situacao, "cancelado")
        self.assertEqual((itens["Garrafa"].situacao, itens["Garrafa"].sub_id), ("pendente", "sem_subid"))

    def test_guarda_o_motivo_so_do_pedido_cancelado(self):
        com_motivo = dict(item("Mop", "60", 1, "5"), displayItemStatus="Cancelled by buyer",
                          itemNotes="Pagamento não aprovado", fraudStatus="FRAUD")
        so_status = dict(item("Fone", "50", 1, "2"), displayItemStatus="Cancelled", itemNotes="Cancelled")
        completo = dict(item("Capa", "10", 1, "1"), displayItemStatus="Completed", itemNotes="obs qualquer")
        self.sincronizar([conversao(ts(2026, 10, 10), [
            pedido("p1", "CANCELLED", [com_motivo, so_status]), pedido("p2", "COMPLETED", [completo]),
        ], conversion_id="c1")])
        itens = {v.nome: v for v in VendaItem.objects.filter(conta=self.conta)}
        self.assertEqual(itens["Mop"].motivo_shopee,
                         "Cancelled by buyer - Pagamento não aprovado - Pedido marcado como fraude pela Shopee")
        self.assertEqual(itens["Fone"].motivo_shopee, "Cancelled")  # texto repetido não duplica
        self.assertEqual(itens["Capa"].motivo_shopee, "")

    def test_ressincronizar_troca_a_janela_e_guarda_o_historico_por_13_meses(self):
        antigo = VendaItem.objects.create(  # fora da janela da API: fica
            conta=self.conta, conversao_id="v", pedido_id="velho", nome="Antigo",
            comprado_em=self.AGORA - timedelta(days=200),
        )
        vencido = VendaItem.objects.create(  # mais de 13 meses: sai
            conta=self.conta, conversao_id="v", pedido_id="vencido", nome="Vencido",
            comprado_em=self.AGORA - timedelta(days=400),
        )
        self.sincronizar([conversao(ts(2026, 10, 10), [pedido("p1", "PENDING", [item("Fone", "50")])])])
        self.sincronizar([conversao(ts(2026, 10, 10), [pedido("p1", "COMPLETED", [item("Fone", "50")])])])
        self.assertEqual(VendaItem.objects.get(nome="Fone").situacao, "concluido")  # sem duplicar
        self.assertTrue(VendaItem.objects.filter(pk=antigo.pk).exists())
        self.assertFalse(VendaItem.objects.filter(pk=vencido.pk).exists())


class TestBuscaDosItens(TestCase):
    def test_tira_campo_do_item_que_a_api_nao_conhece(self):
        respostas = [
            ShopeeAPIError('Cannot query field "shopName" on type "ConversionReportItem".'),
            {"conversionReport": {"nodes": [{"conversionId": "1"}], "pageInfo": {"hasNextPage": False}}},
        ]
        with mock.patch("painel.shopee.executar_graphql_com_credenciais", side_effect=respostas) as executar:
            buscar_conversoes("123", "segredo", 1, 2)
        primeira, segunda = [c.args[2] for c in executar.call_args_list]
        self.assertIn("shopName", primeira)
        self.assertNotIn("shopName", segunda)
        self.assertIn("itemName", segunda)

    def test_campo_essencial_do_item_nao_sai(self):
        erro = ShopeeAPIError('Cannot query field "itemPrice" on type "ConversionReportItem".')
        with mock.patch("painel.shopee.executar_graphql_com_credenciais", side_effect=erro):
            with self.assertRaises(ShopeeAPIError):
                buscar_conversoes("123", "segredo", 1, 2)


class TestTelaDeVendas(BasePainel):
    # Relógio parado no meio do mês: as vendas de "ontem" caem em "Este mês" (no dia 1º, ontem é o
    # mês passado e os testes falhavam).
    AGORA = datetime(2026, 10, 15, 15, 0, tzinfo=SP)

    def setUp(self):
        relogio = mock.patch("django.utils.timezone.now", return_value=self.AGORA)
        relogio.start()
        self.addCleanup(relogio.stop)
        super().setUp()
        self.entrar()
        self.conta.shopee_app_id = "123"
        self.conta.shopee_validado_em = timezone.now()
        self.conta.save()
        self.hoje = timezone.localdate()

    def venda(self, nome, dias_atras=1, situacao="pendente", sub_id="grupoA", comissao="1", qtd=1, item_id=None, **campos):
        momento = timezone.now() - timedelta(days=dias_atras)
        return VendaItem.objects.create(
            conta=self.conta, conversao_id=f"c{VendaItem.objects.count()}", pedido_id=f"p{VendaItem.objects.count()}",
            item_id=item_id or nome, nome=nome, quantidade=qtd, valor=Decimal("10") * qtd, comissao=Decimal(comissao),
            situacao=situacao, sub_id=sub_id, comprado_em=momento, **campos,
        )

    def test_sem_vendas_guardadas_avisa(self):
        html = self.client.get("/resultados/").content.decode()
        self.assertIn("Produtos mais vendidos", html)
        self.assertIn("depois da próxima atualização com a Shopee", html)
        self.assertNotIn("Exportar Excel", html)

    def test_ranking_com_5_sem_cancelados(self):
        for i in range(7):
            self.venda(f"Produto {i}", qtd=i + 1)
        self.venda("Cancelado", qtd=50, situacao="cancelado")
        ranking = self.client.get("/resultados/").context["ranking"]
        self.assertEqual([p["nome"] for p in ranking], ["Produto 6", "Produto 5", "Produto 4", "Produto 3", "Produto 2"])
        self.assertEqual((ranking[0]["unidades"], ranking[0]["barra"], ranking[-1]["barra"]), (7, 100, 43))

    def test_lista_com_20_por_pagina_e_filtros(self):
        RegraSubId.objects.create(conta=self.conta, operador="comeca", texto="user", nome="Site")
        for i in range(23):
            self.venda(f"Fone {i}", sub_id="grupoA", comissao=str(i))
        self.venda("Garrafa", sub_id="user123", situacao="concluido", comissao="99")
        self.venda("Garrafa velha", sub_id="user456", dias_atras=200)  # fora do período

        resposta = self.client.get("/resultados/", {"periodo": "30dias"})
        pagina = resposta.context["pagina_vendas"]
        self.assertEqual((pagina.paginator.count, len(pagina.object_list)), (24, 20))
        self.assertEqual(resposta.context["grupos_vendas"], ["grupoA", "Site"])
        html = resposta.content.decode()
        self.assertIn('name="periodo" value="30dias"', html)  # os filtros mantêm o período
        self.assertIn("Mostrando 1–20 de 24", html)
        self.assertIn("pagina=2", html)

        site = self.client.get("/resultados/", {"periodo": "30dias", "sub": "Site"}).context["pagina_vendas"]
        self.assertEqual([v.nome for v in site], ["Garrafa"])
        self.assertEqual(site[0].grupo, "Site")
        concluidas = self.client.get("/resultados/", {"periodo": "30dias", "situacao": "concluido"}).context["pagina_vendas"]
        self.assertEqual([v.nome for v in concluidas], ["Garrafa"])
        busca = self.client.get("/resultados/", {"periodo": "30dias", "busca": "fone 2"}).context["pagina_vendas"]
        self.assertEqual(sorted(v.nome for v in busca), ["Fone 2", "Fone 20", "Fone 21", "Fone 22"])
        por_comissao = self.client.get("/resultados/", {"periodo": "30dias", "ordem": "comissao"}).context["pagina_vendas"]
        self.assertEqual(por_comissao[0].nome, "Garrafa")

    def test_exportar_excel_com_loja_e_pedido(self):
        self.venda("Fone", loja="Loja do Fone", sub_id="sem_subid")
        self.venda("Mop", situacao="cancelado", motivo_shopee="Cancelado pelo comprador - Pagamento não aprovado")
        resposta = self.client.get("/resultados/vendas.xlsx", {"periodo": "30dias"})
        self.assertEqual(resposta.status_code, 200)
        self.assertIn("attachment;", resposta["Content-Disposition"])
        planilha = load_workbook(io.BytesIO(resposta.content))
        self.assertEqual(planilha.sheetnames, ["Vendas", "Produtos"])
        vendas = list(planilha["Vendas"].values)
        cabecalho = vendas[0]
        self.assertIn("Loja", cabecalho)
        self.assertIn("Pedido", cabecalho)
        linhas = {linha[cabecalho.index("Produto")]: linha for linha in vendas[1:]}
        self.assertEqual(linhas["Fone"][cabecalho.index("Loja")], "Loja do Fone")
        self.assertEqual(linhas["Fone"][cabecalho.index("Grupo")], "Sem Sub ID")
        self.assertEqual(linhas["Mop"][cabecalho.index("Status")], "Cancelado")
        motivo = cabecalho.index("Motivo do cancelamento (Shopee)")
        self.assertEqual(linhas["Mop"][motivo], "Cancelado pelo comprador - Pagamento não aprovado")
        self.assertIn(linhas["Fone"][motivo], (None, ""))  # só o cancelado traz motivo
        produtos = list(planilha["Produtos"].values)
        self.assertEqual([p[0] for p in produtos[1:]], ["Fone"])  # cancelado fora do ranking

    def test_exportar_respeita_os_filtros(self):
        self.venda("Fone", situacao="concluido")
        self.venda("Mop")
        resposta = self.client.get("/resultados/vendas.xlsx", {"periodo": "30dias", "situacao": "concluido"})
        nomes = [linha[2] for linha in list(load_workbook(io.BytesIO(resposta.content))["Vendas"].values)[1:]]
        self.assertEqual(nomes, ["Fone"])


class TestLinksDosFiltros(TestTelaDeVendas):
    def test_links_de_periodo_e_paginas_sem_interrogacao_dupla(self):
        for i in range(25):
            self.venda(f"Fone {i}")
        html = self.client.get("/resultados/", {"periodo": "mes"}).content.decode()
        self.assertNotIn("??", html)
        self.assertIn('href="?periodo=hoje"', html)
        self.assertIn('href="?periodo=mes&amp;pagina=2#vendas"', html)
        self.assertIn('href="/resultados/vendas.xlsx?periodo=mes"', html)


class TestPaginasNoCelular(TestTelaDeVendas):
    def test_celular_mostra_10_por_pagina(self):
        for i in range(25):
            self.venda(f"Fone {i}")
        computador = self.client.get("/resultados/", {"periodo": "30dias"})
        self.assertEqual(len(computador.context["pagina_vendas"].object_list), 20)
        self.assertContains(computador, 'data-tela="computador"')
        self.client.cookies["appf_tela"] = "celular"
        celular = self.client.get("/resultados/", {"periodo": "30dias"})
        self.assertEqual((len(celular.context["pagina_vendas"].object_list), celular.context["pagina_vendas"].paginator.num_pages), (10, 3))
        self.assertContains(celular, 'data-tela="celular"')


class TestSincronizacaoRapidaECompleta(TestCase):
    AGORA = datetime(2026, 10, 18, 15, 0, tzinfo=SP)

    def setUp(self):
        self.conta = ContaAppfiliado.objects.create(
            email="bia@exemplo.com", shopee_app_id="123", shopee_secret_cifrado=cripto.cifrar("segredo"),
            shopee_validado_em=timezone.now(),
        )

    def sincronizar(self, paginas, agora=None, **opcoes):
        with mock.patch("painel.resultados.paginas_de_conversoes", return_value=iter(paginas)) as buscar:
            resumo = resultados.sincronizar_conta(self.conta, agora=agora or self.AGORA, **opcoes)
        return buscar, resumo

    def janela_em_dias(self, buscar, agora=None):
        inicio_ts, fim_ts = buscar.call_args.args[2:]
        return round((fim_ts - inicio_ts) / 86400)

    def test_primeira_completa_depois_rapidas_e_completa_de_madrugada(self):
        buscar, resumo = self.sincronizar([])
        self.assertEqual((self.janela_em_dias(buscar), resumo["completa"]), (89, True))
        buscar, resumo = self.sincronizar([], agora=self.AGORA + timedelta(hours=3))
        self.assertEqual((self.janela_em_dias(buscar), resumo["completa"]), (30, False))
        madrugada = datetime(2026, 10, 19, 4, 0, tzinfo=SP)  # 13 h depois da completa
        buscar, resumo = self.sincronizar([], agora=madrugada)
        self.assertEqual((self.janela_em_dias(buscar), resumo["completa"]), (89, True))
        _, resumo = self.sincronizar([], agora=madrugada + timedelta(hours=3))  # 7h: rápida
        self.assertFalse(resumo["completa"])
        _, resumo = self.sincronizar([], agora=madrugada + timedelta(hours=31))  # passou de 30 h
        self.assertTrue(resumo["completa"])

    def test_rapida_mexe_so_nos_30_dias_e_guarda_o_resto(self):
        compra_antiga = ts(2026, 9, 1)  # 47 dias antes: fora da rápida
        self.sincronizar([[
            conversao(compra_antiga, [pedido("p1", "PENDING", [item("Antigo", "100", 1, "10")])], conversion_id="c1"),
            conversao(ts(2026, 10, 10), [pedido("p2", "PENDING", [item("Novo", "50", 1, "5")])], conversion_id="c2"),
        ]])
        # rápida: a Shopee só devolve o que é dos últimos 30 dias
        self.sincronizar([[
            conversao(ts(2026, 10, 10), [pedido("p2", "COMPLETED", [item("Novo", "50", 1, "5", ts(2026, 10, 17))])], conversion_id="c2"),
        ]], agora=self.AGORA + timedelta(hours=3))
        self.assertEqual(VendaItem.objects.get(nome="Antigo").situacao, "pendente")  # guardado como estava
        self.assertEqual(VendaItem.objects.get(nome="Novo").situacao, "concluido")
        from .models import ResultadoDiario
        from datetime import date
        self.assertEqual(ResultadoDiario.objects.get(conta=self.conta, dia=date(2026, 9, 1)).faturamento, Decimal("100.00"))
        self.assertEqual(ResultadoDiario.objects.get(conta=self.conta, dia=date(2026, 10, 17)).comissao_validada, Decimal("5.00"))

    def test_validada_soma_os_itens_guardados_mesmo_de_compra_antiga(self):
        from datetime import date

        from .models import ResultadoDiario

        # compra de 40 dias atrás validada ontem: só a completa vê, e fica guardada
        self.sincronizar([[
            conversao(ts(2026, 9, 8), [pedido("p1", "COMPLETED", [item("Fone", "100", 1, "12", ts(2026, 10, 17))])]),
        ]])
        self.assertEqual(ResultadoDiario.objects.get(conta=self.conta, dia=date(2026, 10, 17)).comissao_validada, Decimal("12.00"))
        self.sincronizar([], agora=self.AGORA + timedelta(hours=3))  # rápida não vê a compra antiga
        self.assertEqual(ResultadoDiario.objects.get(conta=self.conta, dia=date(2026, 10, 17)).comissao_validada, Decimal("12.00"))

    def test_processa_pagina_por_pagina(self):
        _, resumo = self.sincronizar([
            [conversao(ts(2026, 10, 10), [pedido("p1", "PENDING", [item("A", "10")])], conversion_id="c1")],
            [conversao(ts(2026, 10, 11), [pedido("p2", "PENDING", [item("B", "20")])], conversion_id="c2")],
        ])
        self.assertEqual(resumo["conversoes"], 2)
        self.assertEqual(VendaItem.objects.count(), 2)

    def test_falha_no_meio_nao_apaga_o_que_estava_guardado(self):
        self.sincronizar([[conversao(ts(2026, 10, 10), [pedido("p1", "PENDING", [item("A", "10")])])]])

        def paginas_que_falham(*_):
            yield [conversao(ts(2026, 10, 12), [pedido("p9", "PENDING", [item("B", "10")])], conversion_id="c9")]
            raise ShopeeAPIError("caiu")

        with mock.patch("painel.resultados.paginas_de_conversoes", side_effect=paginas_que_falham):
            with self.assertRaises(ShopeeAPIError):
                resultados.sincronizar_conta(self.conta, agora=self.AGORA + timedelta(hours=3))
        self.assertEqual(list(VendaItem.objects.values_list("nome", flat=True)), ["A"])

    def test_log_mostra_tipo_memoria_e_validacoes_atrasadas(self):
        with mock.patch("painel.resultados.paginas_de_conversoes", return_value=iter([[
            conversao(ts(2026, 9, 1), [pedido("p1", "COMPLETED", [item("A", "10", 1, "1", ts(2026, 10, 15))])], conversion_id="c1"),
            conversao(ts(2026, 10, 1), [pedido("p2", "COMPLETED", [item("B", "10", 1, "1", ts(2026, 10, 10))])], conversion_id="c2"),
        ]])), mock.patch("painel.resultados.timezone.now", return_value=self.AGORA):
            texto = resultados.processar(self.conta)
        self.assertIn("busca completa em ", texto)
        self.assertIn("memória", texto)
        self.assertIn("validadas mais de 30 dias depois da compra: 1 de 2", texto)


class TestAtualizarOntemEValidacoes(TestCase):
    """O botão "Atualizar ontem e validações de hoje": compras de ontem + o que a Shopee validou."""

    AGORA = datetime(2026, 10, 18, 15, 0, tzinfo=SP)

    def setUp(self):
        self.conta = ContaAppfiliado.objects.create(
            email="bia@exemplo.com", shopee_app_id="123", shopee_secret_cifrado=cripto.cifrar("segredo"),
            shopee_validado_em=timezone.now(),
        )
        # Pedido de 05/10 guardado como pendente (a Shopee validou hoje).
        VendaItem.objects.create(conta=self.conta, conversao_id="antiga", pedido_id="P1", posicao=0, item_id="X1", nome="Fone",
                                 valor=Decimal(50), comissao=Decimal("4.00"), situacao=VendaItem.PENDENTE, sub_id="grupoA",
                                 comprado_em=datetime(2026, 10, 5, 12, tzinfo=SP))
        from .models import SincronizacaoResultados
        SincronizacaoResultados.objects.create(conta=self.conta, vendas_desde=datetime(2026, 7, 20, tzinfo=SP))

    def test_busca_as_compras_de_ontem_e_as_validacoes_de_compras_antigas(self):
        validado_hoje = ts(2026, 10, 18, 9)

        def item_(item_id, preco, comissao, validado=None):
            return {"itemId": item_id, "itemName": item_id, "itemPrice": preco, "qty": 1, "itemTotalCommission": comissao,
                    "completeTime": validado}

        de_ontem = conversao(ts(2026, 10, 17), [pedido("P9", "PENDING", [item_("Y1", "30", "2")])], conversion_id="nova")
        validadas = [
            conversao(ts(2026, 10, 5), [pedido("P1", "COMPLETED", [item_("X1", "50", "4.20", validado_hoje)])], conversion_id="antiga"),
            conversao(ts(2026, 9, 28), [pedido("P2", "COMPLETED", [item_("Z1", "80", "6", validado_hoje)])], conversion_id="nao-guardada"),
            de_ontem,  # comprada ontem: já vem na busca pela compra
        ]

        def paginas(app_id, secret, inicio, fim, data="purchaseTime", itens_extras=()):
            return iter([validadas] if data == "completeTime" else [[de_ontem]])

        with mock.patch("painel.resultados.paginas_de_conversoes", side_effect=paginas) as buscar:
            resumo = resultados.sincronizar_conta(self.conta, agora=self.AGORA, curta=True)
        self.assertEqual([c.kwargs.get("data", "purchaseTime") for c in buscar.call_args_list], ["purchaseTime", "completeTime"])
        self.assertEqual(resumo["validacoes"], 2)
        antiga = VendaItem.objects.get(conversao_id="antiga")
        self.assertEqual((antiga.situacao, antiga.comissao, antiga.sub_id), (VendaItem.CONCLUIDO, Decimal("4.20"), "grupoA"))
        self.assertEqual(timezone.localtime(antiga.concluido_em).date(), date(2026, 10, 18))
        self.assertTrue(VendaItem.objects.filter(conversao_id="nao-guardada", situacao=VendaItem.CONCLUIDO).exists())
        self.assertEqual(VendaItem.objects.filter(conversao_id="nova").count(), 1)  # sem duplicar
        from .models import ResultadoDiario
        self.assertEqual(ResultadoDiario.objects.get(conta=self.conta, dia=date(2026, 10, 18)).comissao_validada, Decimal("10.20"))
